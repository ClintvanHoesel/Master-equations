"""Named simulation results, portable numeric storage, and tabular exports."""

import csv
import json
from dataclasses import dataclass, field
from pathlib import Path

import numpy as np

from .basis import RadialBasis
from .config import Numerics, Parameters


@dataclass(frozen=True)
class Result:
    time: np.ndarray
    state: np.ndarray
    mean_field: np.ndarray
    parameters: Parameters
    numerics: Numerics
    basis: RadialBasis
    evaluations: int
    converged: bool = False
    residual: float | None = None
    solver_settings: dict = field(default_factory=dict)

    def __post_init__(self):
        for name in ("time", "state", "mean_field"):
            value = np.array(getattr(self, name), dtype=float, copy=True)
            if not np.all(np.isfinite(value)):
                raise ValueError(f"non-finite {name} in result")
            value.setflags(write=False)
            object.__setattr__(self, name, value)
        if (
            self.time.ndim != 1
            or self.time.size == 0
            or self.time[0] < 0
            or np.any(np.diff(self.time) <= 0)
        ):
            raise ValueError("result times must be nonnegative and strictly increasing")
        if self.state.shape != (self.time.size, 4 + 2 * self.basis.size):
            raise ValueError(
                "result state shape does not match the time and radial grids"
            )
        if self.mean_field.shape != self.time.shape:
            raise ValueError("mean-field shape does not match time")

    @property
    def density(self):
        return self.state[:, 3]

    @property
    def losses(self):
        """Cumulative lost densities: radiative decay, TTA, TPQ."""
        return self.state[:, :3]

    @property
    def g_tt(self):
        return self.state[:, 4 : 4 + self.basis.size] / self.basis.uniform

    @property
    def g_tp(self):
        return self.state[:, 4 + self.basis.size :] / self.basis.uniform

    @property
    def effective_rates(self):
        p = self.parameters
        edges = self.basis.edges
        integral = (edges[:-1] ** -3 - edges[1:] ** -3) / 3
        return (
            4
            * np.pi
            * p.decay_rate
            * np.column_stack(
                (
                    p.tta_radius**6 * (self.g_tt @ integral),
                    p.tpq_radius**6 * (self.g_tp @ integral),
                )
            )
        )

    @property
    def loss_rates(self):
        ktt, ktp = self.effective_rates.T
        return np.column_stack(
            (
                self.parameters.decay_rate * self.density,
                self.density**2 * ktt,
                self.density * self.parameters.polaron_density * ktp,
            )
        )

    @property
    def loss_fractions(self):
        total = self.losses.sum(axis=1, keepdims=True)
        return np.divide(
            self.losses, total, out=np.zeros_like(self.losses), where=total > 0
        )

    @property
    def metadata(self):
        return dict(
            format_version=1,
            parameters=self.parameters.to_dict(),
            numerics=self.numerics.to_dict(),
            evaluations=self.evaluations,
            converged=self.converged,
            residual=self.residual,
            solver_settings=self.solver_settings,
            units=dict(
                time="microsecond",
                radius="nm",
                density="nm^-3",
                effective_rate="nm^3/microsecond",
            ),
        )

    def summary(self):
        rates = self.effective_rates[-1]
        losses = self.loss_rates[-1]
        return dict(
            final_time=float(self.time[-1]),
            density=float(self.density[-1]),
            k_tta=float(rates[0]),
            k_tpq=float(rates[1]),
            decay_loss_rate=float(losses[0]),
            tta_loss_rate=float(losses[1]),
            tpq_loss_rate=float(losses[2]),
            converged=self.converged,
            residual=self.residual,
            evaluations=self.evaluations,
        )

    def save(self, path: str | Path):
        """Save complete state and configuration without executable pickle data."""
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        with path.open("wb") as handle:
            np.savez_compressed(
                handle,
                time=self.time,
                state=self.state,
                mean_field=self.mean_field,
                edges=self.basis.edges,
                metadata=json.dumps(self.metadata, sort_keys=True),
            )
        return path

    @classmethod
    def load(cls, path: str | Path):
        with np.load(path, allow_pickle=False) as data:
            meta = json.loads(str(data["metadata"]))
            if meta["format_version"] != 1:
                raise ValueError("unsupported result format version")
            return cls(
                data["time"],
                data["state"],
                data["mean_field"],
                Parameters(**meta["parameters"]),
                Numerics(**meta["numerics"]),
                RadialBasis(data["edges"]),
                meta["evaluations"],
                meta["converged"],
                meta["residual"],
                meta.get("solver_settings", {}),
            )

    def _table(self):
        header = [
            "time_us",
            "mean_field_nm^-3",
            "density_nm^-3",
            "decay_loss_nm^-3",
            "tta_loss_nm^-3",
            "tpq_loss_nm^-3",
            "k_tta_nm^3_per_us",
            "k_tpq_nm^3_per_us",
        ]
        return header, np.column_stack(
            (
                self.time,
                self.mean_field,
                self.density,
                self.losses,
                self.effective_rates,
            )
        )

    def to_csv(self, path: str | Path):
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        header, values = self._table()
        with path.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.writer(handle)
            writer.writerow(header)
            writer.writerows(values)
        return path

    def to_excel(self, path: str | Path):
        from openpyxl import Workbook

        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook(write_only=True)
        sheet = workbook.create_sheet("Decay Numbers")
        header, values = self._table()
        sheet.append(header)
        for row in values:
            sheet.append(row.tolist())
        for name, correlation in (
            ("NormalisedG2TT", self.g_tt),
            ("NormalisedG2TP", self.g_tp),
        ):
            sheet = workbook.create_sheet(name)
            sheet.append(["time_us"] + [f"r={r:.10g} nm" for r in self.basis.centers])
            for time, row in zip(self.time, correlation, strict=True):
                sheet.append([float(time), *row.tolist()])
        sheet = workbook.create_sheet("Final Correlation")
        sheet.append(["radius_nm", "g_TT", "g_TP"])
        for row in zip(self.basis.centers, self.g_tt[-1], self.g_tp[-1], strict=True):
            sheet.append(list(row))
        sheet = workbook.create_sheet("Loss fractions")
        sheet.append(["time_us", "decay", "TTA", "TPQ"])
        for time, row in zip(self.time, self.loss_fractions, strict=True):
            sheet.append([float(time), *row.tolist()])
        sheet = workbook.create_sheet("Metadata")
        sheet.append(["key", "value"])
        for key, value in self.metadata.items():
            sheet.append([key, json.dumps(value, sort_keys=True)])
        workbook.save(path)
        return path
